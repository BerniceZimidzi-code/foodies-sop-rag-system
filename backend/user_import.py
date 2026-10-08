import io
import re
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


MAX_WORKBOOK_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_WORKBOOK_ENTRIES = 2000
MAX_IMPORT_ROWS = 500
REQUIRED_COLUMNS = ("name", "email", "department", "role")
ALLOWED_ROLES = {"staff", "manager", "admin"}


class UserImportError(ValueError):
    pass


def parse_user_workbook(filename: str, content: bytes) -> list[dict[str, Any]]:
    if Path(filename).suffix.casefold() != ".xlsx":
        raise UserImportError("Upload an Excel workbook in .xlsx format.")
    if not content:
        raise UserImportError("The spreadsheet file is empty.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise UserImportError("The spreadsheet exceeds the 5 MB size limit.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_WORKBOOK_ENTRIES
                or sum(entry.file_size for entry in entries)
                > MAX_EXPANDED_WORKBOOK_BYTES
            ):
                raise UserImportError("The spreadsheet expands beyond the allowed size.")
    except (zipfile.BadZipFile, OSError) as exc:
        raise UserImportError("The spreadsheet is invalid or could not be read.") from exc
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise UserImportError("The spreadsheet is invalid or could not be read.") from exc

    try:
        worksheet = workbook.active
        if worksheet.max_row is not None and worksheet.max_row > MAX_IMPORT_ROWS + 1:
            raise UserImportError(
                f"The spreadsheet may contain no more than {MAX_IMPORT_ROWS} user rows."
            )
        rows = worksheet.iter_rows(values_only=True)
        header_values = next(rows, None)
        if not header_values:
            raise UserImportError("The spreadsheet must include a header row.")
        headers = [
            str(value).strip().casefold() if value is not None else ""
            for value in header_values
        ]
        if len(headers) != len(set(headers)):
            raise UserImportError("The spreadsheet has duplicate column headers.")
        missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing_columns:
            raise UserImportError(
                "Missing required columns: " + ", ".join(missing_columns) + "."
            )
        positions = {column: headers.index(column) for column in REQUIRED_COLUMNS}
        parsed: list[dict[str, Any]] = []
        seen_emails: set[str] = set()
        for row_number, values in enumerate(rows, start=2):
            if all(value is None or not str(value).strip() for value in values):
                continue
            if len(parsed) >= MAX_IMPORT_ROWS:
                raise UserImportError(
                    f"The spreadsheet may contain no more than {MAX_IMPORT_ROWS} user rows."
                )
            record = {
                column: (
                    str(values[position]).strip()
                    if position < len(values) and values[position] is not None
                    else ""
                )
                for column, position in positions.items()
            }
            record["role"] = record["role"].casefold()
            errors = validate_import_row(record)
            normalized_email = record["email"].casefold()
            if normalized_email and normalized_email in seen_emails:
                errors.append("Email is duplicated in this spreadsheet.")
            if normalized_email:
                seen_emails.add(normalized_email)
            parsed.append(
                {
                    "row": row_number,
                    **record,
                    "errors": errors,
                    "valid": not errors,
                }
            )
        return parsed
    finally:
        workbook.close()


def validate_import_row(record: dict[str, Any]) -> list[str]:
    errors = []
    name = record.get("name", "").strip()
    email = record.get("email", "").strip().casefold()
    department = record.get("department", "").strip()
    role = record.get("role", "").strip().casefold()
    if not name:
        errors.append("Name is required.")
    elif len(name) > 120:
        errors.append("Name must be 120 characters or fewer.")
    if not email or len(email) > 254 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        errors.append("A valid email address is required.")
    if not department:
        errors.append("Department is required.")
    elif len(department) > 100:
        errors.append("Department must be 100 characters or fewer.")
    if role not in ALLOWED_ROLES:
        errors.append("Role must be staff, manager, or admin.")
    return errors
